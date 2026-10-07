"""Local-only calibration register. No external network calls or services."""

from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import os
import re
import sqlite3
import threading
import webbrowser
import zipfile
from zipfile import BadZipFile
import sys
from calendar import monthrange
from contextlib import closing
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from openpyxl import Workbook, load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parent
DATA = (Path(os.environ.get("LOCALAPPDATA", Path.home())) / "CNAS计量台账"
        if getattr(sys, "frozen", False) else ROOT / "data")
PDF_DIR = DATA / "certificates"
DB_FILE = DATA / "register.sqlite3"
MAX_UPLOAD = 15 * 1024 * 1024
BACKUP_LOCK = threading.Lock()
EQUIPMENT_FIELDS = (
    "asset_code", "name", "model", "serial_no", "manufacturer", "category",
    "department", "location", "calibration_cycle", "stopped", "abnormal_flag",
    "abnormal_note",
)
HEADERS = {
    "资产编号": "asset_code", "设备编号": "asset_code", "名称": "name", "设备名称": "name",
    "型号": "model", "出厂编号": "serial_no", "生产厂家": "manufacturer",
    "分类": "category", "ABC分类": "category", "使用部门": "department",
    "存放地点": "location", "检定周期(月)": "calibration_cycle",
    "检定周期": "calibration_cycle", "约定金额": "agreed_amount",
    "证书编号": "cert_no", "金额": "amount", "校准日期": "cal_date",
}


def now():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def iso_date(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip().replace("年", "-").replace("月", "-").replace("日", "")
    text = re.sub(r"[./]", "-", text)
    match = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if not match:
        raise ValueError(f"日期格式应为 YYYY-MM-DD：{value}")
    return date(*map(int, match.groups())).isoformat()


def add_months(day, months):
    y = day.year + (day.month - 1 + months) // 12
    m = (day.month - 1 + months) % 12 + 1
    return day.replace(year=y, month=m, day=min(day.day, monthrange(y, m)[1]))


def amount(value):
    if value is None or str(value).strip() == "":
        return None
    try:
        number = Decimal(str(value).replace(",", "").replace("￥", "").strip())
    except InvalidOperation as exc:
        raise ValueError("金额格式不正确") from exc
    if not number.is_finite() or number < 0:
        raise ValueError("金额不能为负数")
    return str(number.quantize(Decimal("0.01")))


def connect(db_path=DB_FILE):
    db_path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(db_path, timeout=15)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys=ON")
    connection.execute("PRAGMA journal_mode=WAL")
    return connection


def init_db(db_path=DB_FILE):
    with closing(connect(db_path)) as db:
        db.executescript("""
        CREATE TABLE IF NOT EXISTS equipment (
          id INTEGER PRIMARY KEY, asset_code TEXT NOT NULL UNIQUE, name TEXT NOT NULL,
          model TEXT NOT NULL DEFAULT '', serial_no TEXT NOT NULL DEFAULT '',
          manufacturer TEXT NOT NULL DEFAULT '', category TEXT NOT NULL DEFAULT 'C'
            CHECK(category IN ('A','B','C')),
          department TEXT NOT NULL DEFAULT '', location TEXT NOT NULL DEFAULT '',
          calibration_cycle INTEGER NOT NULL DEFAULT 12 CHECK(calibration_cycle BETWEEN 1 AND 120),
          stopped INTEGER NOT NULL DEFAULT 0, deleted INTEGER NOT NULL DEFAULT 0,
          abnormal_flag TEXT NOT NULL DEFAULT '',
          abnormal_note TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL, updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS certificate (
          id INTEGER PRIMARY KEY, equipment_id INTEGER REFERENCES equipment(id) ON DELETE SET NULL,
          cert_no TEXT NOT NULL DEFAULT '', file_name TEXT NOT NULL, file_path TEXT NOT NULL,
          cert_hash TEXT NOT NULL UNIQUE, cal_org TEXT NOT NULL DEFAULT '',
          cal_date TEXT, valid_until TEXT, result TEXT NOT NULL DEFAULT '',
          agreed_amount TEXT, state TEXT NOT NULL DEFAULT 'pending'
            CHECK(state IN ('pending','confirmed','rejected')),
          review_note TEXT NOT NULL DEFAULT '', parsed_data TEXT NOT NULL DEFAULT '{}',
          created_at TEXT NOT NULL, reviewed_at TEXT
        );
        CREATE INDEX IF NOT EXISTS idx_cert_equipment ON certificate(equipment_id,cal_date);
        CREATE TABLE IF NOT EXISTS billing_record (
          id INTEGER PRIMARY KEY, batch_no TEXT NOT NULL, line_no INTEGER NOT NULL,
          asset_code TEXT NOT NULL DEFAULT '', cert_no TEXT NOT NULL DEFAULT '',
          amount TEXT, raw_data TEXT NOT NULL, finding TEXT NOT NULL,
          detail TEXT NOT NULL, imported_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_billing_batch ON billing_record(batch_no);
        CREATE TABLE IF NOT EXISTS submission (
          id INTEGER PRIMARY KEY, equipment_id INTEGER NOT NULL REFERENCES equipment(id),
          sent_date TEXT NOT NULL, returned_date TEXT, note TEXT NOT NULL DEFAULT '',
          created_at TEXT NOT NULL, updated_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_submission_date ON submission(sent_date);
        CREATE TABLE IF NOT EXISTS audit_log (
          id INTEGER PRIMARY KEY, entity TEXT NOT NULL, entity_id INTEGER,
          action TEXT NOT NULL, old_data TEXT, new_data TEXT, operator TEXT NOT NULL,
          timestamp TEXT NOT NULL
        );
        """)
        if "deleted" not in {row[1] for row in db.execute("PRAGMA table_info(equipment)")}:
            db.execute("ALTER TABLE equipment ADD COLUMN deleted INTEGER NOT NULL DEFAULT 0")
        db.commit()


def audit(db, entity, entity_id, action, old, new, operator):
    db.execute("INSERT INTO audit_log(entity,entity_id,action,old_data,new_data,operator,timestamp) VALUES(?,?,?,?,?,?,?)",
               (entity, entity_id, action, json.dumps(old, ensure_ascii=False, default=str),
                json.dumps(new, ensure_ascii=False, default=str), operator or "本机用户", now()))


def row_dict(row):
    return dict(row) if row else None


def active_cert(db, equipment_id):
    return db.execute("""SELECT * FROM certificate WHERE equipment_id=? AND state='confirmed'
        AND result='合格' AND cal_date IS NOT NULL ORDER BY cal_date DESC,id DESC LIMIT 1""",
        (equipment_id,)).fetchone()


def latest_cert(db, equipment_id):
    return db.execute("""SELECT * FROM certificate WHERE equipment_id=? AND state='confirmed'
        AND cal_date IS NOT NULL ORDER BY cal_date DESC,id DESC LIMIT 1""",
        (equipment_id,)).fetchone()


def equipment_view(db, row):
    item = dict(row)
    cert = active_cert(db, item["id"])
    latest = latest_cert(db, item["id"])
    item["last_cal_date"] = cert["cal_date"] if cert else None
    item["next_cal_date"] = (cert["valid_until"] or add_months(
        date.fromisoformat(cert["cal_date"]), item["calibration_cycle"]).isoformat()) if cert else None
    if item["stopped"]:
        item["status"] = "停用"
    elif item["abnormal_flag"] or (latest and latest["result"] == "不合格"):
        item["status"] = "异常"
    elif not item["next_cal_date"]:
        item["status"] = "待校准"
    elif item["next_cal_date"] <= date.today().isoformat():
        item["status"] = "超期"
    else:
        item["status"] = "正常"
    item["certificate_issue"] = (f"最新证书 {latest['cert_no']} 不合格"
        if latest and latest["result"] == "不合格" else "")
    return item


def archive_equipment(db, equipment_id, note, operator):
    old = row_dict(db.execute("SELECT * FROM equipment WHERE id=? AND deleted=0", (equipment_id,)).fetchone())
    if not old:
        raise ValueError("设备不存在")
    if not note.strip():
        raise ValueError("归档必须填写原因")
    db.execute("UPDATE equipment SET deleted=1,updated_at=? WHERE id=?", (now(), equipment_id))
    new = row_dict(db.execute("SELECT * FROM equipment WHERE id=?", (equipment_id,)).fetchone())
    audit(db, "equipment", equipment_id, "归档：" + note.strip(), old, new, operator)
    return {"id": equipment_id, "archived": True}


def save_submission(db, data, operator):
    code = str(data.get("asset_code") or "").strip()
    device = db.execute("SELECT id FROM equipment WHERE asset_code=? AND deleted=0", (code,)).fetchone()
    if not device:
        raise ValueError("资产编号未在台账中找到")
    sent = iso_date(data.get("sent_date"))
    returned = iso_date(data.get("returned_date"))
    if not sent or sent > date.today().isoformat():
        raise ValueError("送检日期不能为空或晚于今天")
    if returned and (returned < sent or returned > date.today().isoformat()):
        raise ValueError("返回日期必须在送检日期和今天之间")
    cursor = db.execute("""INSERT INTO submission(equipment_id,sent_date,returned_date,note,created_at,updated_at)
        VALUES(?,?,?,?,?,?)""", (device["id"], sent, returned, str(data.get("note") or "").strip(), now(), now()))
    result = row_dict(db.execute("SELECT * FROM submission WHERE id=?", (cursor.lastrowid,)).fetchone())
    audit(db, "submission", cursor.lastrowid, "送检登记", None, result, operator)
    return result


def return_submission(db, submission_id, data, operator):
    old = row_dict(db.execute("SELECT * FROM submission WHERE id=?", (submission_id,)).fetchone())
    if not old:
        raise ValueError("送检记录不存在")
    if old["returned_date"]:
        raise ValueError("该记录已登记返回")
    returned = iso_date(data.get("returned_date"))
    if not returned or returned < old["sent_date"] or returned > date.today().isoformat():
        raise ValueError("返回日期必须在送检日期和今天之间")
    db.execute("UPDATE submission SET returned_date=?,updated_at=? WHERE id=?", (returned, now(), submission_id))
    result = row_dict(db.execute("SELECT * FROM submission WHERE id=?", (submission_id,)).fetchone())
    audit(db, "submission", submission_id, "返回登记", old, result, operator)
    return result


def equipment_payload(data, old=None):
    result = dict(old or {})
    for field in EQUIPMENT_FIELDS:
        if field in data:
            result[field] = data[field]
    result["asset_code"] = str(result.get("asset_code", "")).strip()
    result["name"] = str(result.get("name", "")).strip()
    if not result["asset_code"] or not result["name"]:
        raise ValueError("资产编号和设备名称必填")
    for field in ("model", "serial_no", "manufacturer", "department", "location", "abnormal_flag", "abnormal_note"):
        result[field] = str(result.get(field) or "").strip()
    result["category"] = str(result.get("category") or "C").upper()
    if result["category"] not in ("A", "B", "C"):
        raise ValueError("ABC 分类只能是 A/B/C")
    result["calibration_cycle"] = int(result.get("calibration_cycle") or 12)
    if not 1 <= result["calibration_cycle"] <= 120:
        raise ValueError("校准周期应在 1 到 120 个月之间")
    result["stopped"] = int(str(result.get("stopped", 0)).lower() in ("1", "true", "yes"))
    if result["abnormal_flag"] not in ("", "计量不过", "异常未检", "其他"):
        raise ValueError("异常类型无效")
    if result["abnormal_flag"] and not result["abnormal_note"]:
        raise ValueError("异常必须填写备注")
    return {field: result[field] for field in EQUIPMENT_FIELDS}


def save_equipment(db, data, operator, equipment_id=None):
    old_row = db.execute("SELECT * FROM equipment WHERE id=? AND deleted=0", (equipment_id,)).fetchone() if equipment_id else None
    if equipment_id and not old_row:
        raise ValueError("设备不存在")
    old = dict(old_row) if old_row else None
    payload = equipment_payload(data, old)
    closure_note = str(data.get("closure_note") or "").strip()
    if old and old["abnormal_flag"] and not payload["abnormal_flag"]:
        if not closure_note:
            raise ValueError("解除异常必须填写处理原因")
        payload["abnormal_note"] = f"{old['abnormal_note']} | 解除：{closure_note}"
    stamp = now()
    if old:
        fields = list(EQUIPMENT_FIELDS)
        db.execute(f"UPDATE equipment SET {','.join(f'{field}=?' for field in fields)},updated_at=? WHERE id=?",
                   [payload[field] for field in fields] + [stamp, equipment_id])
        action = "异常解除" if closure_note and old["abnormal_flag"] and not payload["abnormal_flag"] else "修改"
    else:
        fields = list(EQUIPMENT_FIELDS)
        cursor = db.execute(f"INSERT INTO equipment({','.join(fields)},created_at,updated_at) VALUES({','.join('?' for _ in fields)},?,?)",
                            [payload[field] for field in fields] + [stamp, stamp])
        equipment_id = cursor.lastrowid
        action = "新增"
    new = row_dict(db.execute("SELECT * FROM equipment WHERE id=?", (equipment_id,)).fetchone())
    audit(db, "equipment", equipment_id, action, old, {**new, "closure_note": closure_note} if action == "异常解除" else new, operator)
    return equipment_view(db, new)


def extract_pdf(raw):
    try:
        reader = PdfReader(io.BytesIO(raw), strict=False)
        if reader.is_encrypted and reader.decrypt("") == 0:
            return {}, "PDF 有打开密码，请人工处理"
        text = "\n".join((page.extract_text() or "") for page in reader.pages[:8])
    except Exception as exc:
        raise ValueError("不是可读取的 PDF 文件") from exc
    if not text.strip():
        return {}, "PDF 无可提取文本，请人工录入"
    return parse_certificate_text(text)


def parse_certificate_text(text):
    patterns = {
        "asset_code": r"(?:资产编号|设备编号|管理编号|Asset\s*(?:Code|ID)|Equipment\s*(?:Code|ID))[ \t]*[:：]?[ \t]*([A-Za-z0-9_-]{2,64})",
        "cert_no": r"(?:证书编号|证书号|Certificate\s*(?:No\.?|Number))[ \t]*[:：]?[ \t]*([A-Za-z0-9_/-]{2,64})",
        "cal_date": r"(?:校准日期|检定日期|Calibration\s*Date)[ \t]*[:：]?[ \t]*(\d{4}[./年-]\d{1,2}[./月-]\d{1,2})",
        "valid_until": r"(?:有效期至|下次校准日期|有效日期|Valid\s*Until)[ \t]*[:：]?[ \t]*(\d{4}[./年-]\d{1,2}[./月-]\d{1,2})",
        "cal_org": r"(?:校准机构|检定机构|出具机构|Calibration\s*Organization)[ \t]*[:：]?[ \t]*([^\n]{2,80})",
    }
    found = {}
    for field, pattern in patterns.items():
        matches = re.findall(pattern, text, flags=re.IGNORECASE)
        if len(set(matches)) == 1:
            found[field] = matches[0].strip()
    for field in ("cal_date", "valid_until"):
        if field in found:
            try:
                found[field] = iso_date(found[field])
            except ValueError:
                del found[field]
    result = re.findall(r"(?:校准结论|检定结论|校准结果|检定结果|Calibration\s*Result)[ \t]*[:：]?[ \t]*(不合格|合格|Fail|Pass)", text, flags=re.IGNORECASE)
    normalized = {"pass": "合格", "fail": "不合格"}
    conclusions = {normalized.get(item.lower(), item) for item in result}
    if len(conclusions) == 1:
        found["result"] = conclusions.pop()
    found["text_preview"] = text[:2500]
    return found, "请核对原件后确认"


def upload_pdf(db, data, operator, pdf_dir=PDF_DIR):
    name = str(data.get("file_name") or "").strip()
    if not name.lower().endswith(".pdf"):
        raise ValueError("只能上传 PDF")
    try:
        raw = base64.b64decode(data["content"], validate=True)
    except Exception as exc:
        raise ValueError("文件编码无效") from exc
    if len(raw) > MAX_UPLOAD or not raw.startswith(b"%PDF-"):
        raise ValueError("PDF 无效或超过 15MB")
    digest = hashlib.sha256(raw).hexdigest()
    if db.execute("SELECT id FROM certificate WHERE cert_hash=?", (digest,)).fetchone():
        raise ValueError("此 PDF 已上传")
    found, note = extract_pdf(raw)
    equipment = db.execute("SELECT id FROM equipment WHERE asset_code=? AND deleted=0", (found.get("asset_code", ""),)).fetchone()
    pdf_dir.mkdir(parents=True, exist_ok=True)
    target = pdf_dir / f"{digest}.pdf"
    target.write_bytes(raw)
    try:
        cursor = db.execute("""INSERT INTO certificate(equipment_id,cert_no,file_name,file_path,cert_hash,cal_org,
            cal_date,valid_until,result,parsed_data,review_note,created_at)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
            (equipment["id"] if equipment else None, found.get("cert_no", ""), name,
             str(target), digest, found.get("cal_org", ""), found.get("cal_date"),
             found.get("valid_until"), found.get("result", ""), json.dumps(found, ensure_ascii=False), note, now()))
        certificate_id = cursor.lastrowid
        audit(db, "certificate", certificate_id, "上传", None, found, operator)
        auto_ready = (equipment is not None and found.get("cert_no") and found.get("cal_date")
            and found.get("result") in ("合格", "不合格")
            and found["cal_date"] <= date.today().isoformat()
            and (not found.get("valid_until") or found["valid_until"] > found["cal_date"])
            and not db.execute("SELECT 1 FROM certificate WHERE cert_no=? AND state='confirmed' AND id<>?",
                (found["cert_no"], certificate_id)).fetchone())
        if auto_ready:
            review_certificate(db, certificate_id, {"state": "confirmed", "review_note": "字段完整且设备编号精确匹配，自动确认"}, "自动解析")
        else:
            db.execute("UPDATE certificate SET review_note=? WHERE id=?",
                ("字段不完整、设备未匹配或数据冲突，请核对原件", certificate_id))
        return row_dict(db.execute("SELECT * FROM certificate WHERE id=?", (certificate_id,)).fetchone())
    except Exception:
        target.unlink(missing_ok=True)
        raise


def review_certificate(db, certificate_id, data, operator):
    old = row_dict(db.execute("SELECT * FROM certificate WHERE id=?", (certificate_id,)).fetchone())
    if not old:
        raise ValueError("证书不存在")
    if old["state"] == "rejected":
        raise ValueError("已驳回证书不可再次确认")
    state = data.get("state", "confirmed")
    if state not in ("confirmed", "rejected"):
        raise ValueError("审核状态无效")
    code = str(data.get("asset_code") or "").strip()
    if code:
        matched = db.execute("SELECT id FROM equipment WHERE asset_code=? AND deleted=0", (code,)).fetchone()
        equipment_id = matched["id"] if matched else 0
    else:
        equipment_id = int(data.get("equipment_id") or old["equipment_id"] or 0)
    if state == "confirmed" and not db.execute("SELECT 1 FROM equipment WHERE id=? AND deleted=0", (equipment_id,)).fetchone():
        raise ValueError("必须关联已有设备")
    cert_no = str(data.get("cert_no", old["cert_no"])).strip()
    cal_date = iso_date(data.get("cal_date", old["cal_date"]))
    valid_until = iso_date(data.get("valid_until", old["valid_until"]))
    result = str(data.get("result", old["result"])).strip()
    if state == "confirmed":
        if not cert_no or not cal_date or result not in ("合格", "不合格"):
            raise ValueError("确认需要设备、证书编号、校准日期和明确结果")
        if cal_date > date.today().isoformat():
            raise ValueError("校准日期不能在未来")
        if valid_until and valid_until <= cal_date:
            raise ValueError("有效期必须晚于校准日期")
        duplicate = db.execute("SELECT id FROM certificate WHERE cert_no=? AND state='confirmed' AND id<>?",
                               (cert_no, certificate_id)).fetchone()
        if duplicate:
            raise ValueError("该证书编号已有确认记录")
    agreed = amount(data.get("agreed_amount", old["agreed_amount"]))
    db.execute("""UPDATE certificate SET equipment_id=?,cert_no=?,cal_org=?,cal_date=?,valid_until=?,
        result=?,agreed_amount=?,state=?,review_note=?,reviewed_at=? WHERE id=?""",
        (equipment_id or None, cert_no, str(data.get("cal_org", old["cal_org"])), cal_date,
         valid_until, result, agreed, state, str(data.get("review_note", "")), now(), certificate_id))
    new = row_dict(db.execute("SELECT * FROM certificate WHERE id=?", (certificate_id,)).fetchone())
    audit(db, "certificate", certificate_id, "审核", old, new, operator)
    return new


def sheet_rows(raw, file_name):
    if len(raw) > MAX_UPLOAD:
        raise ValueError("文件超过 15MB")
    if file_name.lower().endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except (BadZipFile, ValueError, KeyError) as exc:
            raise ValueError("Excel 文件无法读取") from exc
        try:
            values = iter(wb.active.values)
            header = next(values, None)
            if not header:
                raise ValueError("表格为空")
            for line_no, values_row in enumerate(values, 2):
                if any(value is not None and str(value).strip() for value in values_row):
                    yield line_no, dict(zip([str(v).strip() if v is not None else "" for v in header], values_row))
        finally:
            wb.close()
    elif file_name.lower().endswith(".csv"):
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV 请保存为 UTF-8 编码") from exc
        for line_no, row in enumerate(csv.DictReader(io.StringIO(text)), 2):
            if any(str(v or "").strip() for v in row.values()):
                yield line_no, row
    else:
        raise ValueError("仅支持 .xlsx 或 UTF-8 CSV")


def decoded_upload(data):
    try:
        raw = base64.b64decode(data["content"], validate=True)
    except Exception as exc:
        raise ValueError("文件编码无效") from exc
    return raw, str(data.get("file_name") or "")


def normalized_row(row):
    return {HEADERS.get(str(key).strip(), str(key).strip()): value for key, value in row.items()}


def import_equipment(db, data, operator):
    raw, name = decoded_upload(data)
    rows = list(sheet_rows(raw, name))
    if len(rows) > 20000:
        raise ValueError("每次最多导入 20000 行")
    success, errors = 0, []
    for line_no, original in rows:
        row = normalized_row(original)
        code = str(row.get("asset_code") or "").strip()
        if not code:
            errors.append({"line": line_no, "error": "缺少资产编号"})
            continue
        try:
            with db:
                existing = db.execute("SELECT id FROM equipment WHERE asset_code=?", (code,)).fetchone()
                save_equipment(db, row, operator, existing["id"] if existing else None)
            success += 1
        except (ValueError, sqlite3.IntegrityError) as exc:
            errors.append({"line": line_no, "error": str(exc)})
    return {"imported": success, "errors": errors}


def reconcile_line(db, code, cert_no, charged):
    if not code and not cert_no:
        return "缺少标识", "需提供资产编号或证书编号"
    equipment = db.execute("SELECT id FROM equipment WHERE asset_code=?", (code,)).fetchone() if code else None
    cert = db.execute("SELECT * FROM certificate WHERE cert_no=? AND state='confirmed'", (cert_no,)).fetchone() if cert_no else None
    if code and not equipment:
        return "设备缺失", "资产编号未在台账中找到；不自动模糊关联"
    if cert_no and not cert:
        return "证书缺失", "未找到已确认的证书"
    if cert and equipment and cert["equipment_id"] != equipment["id"]:
        return "关联不匹配", "证书对应另一台设备"
    if not cert:
        return "待核对", "缺少证书编号，无法确认计量服务"
    if charged is not None and cert["agreed_amount"] is not None:
        if abs(Decimal(charged) - Decimal(cert["agreed_amount"])) > Decimal("0.01"):
            return "金额差异", f"清单 {charged}，约定 {cert['agreed_amount']}"
    if charged is not None and cert["agreed_amount"] is None:
        return "金额待核对", "未录入此证书对应服务的约定金额"
    return "匹配", "设备、证书精确匹配"


def import_billing(db, data, operator):
    raw, name = decoded_upload(data)
    rows = list(sheet_rows(raw, name))
    if len(rows) > 20000:
        raise ValueError("每次最多导入 20000 行")
    period = str(data.get("calibration_month") or "").strip()
    if period and not re.fullmatch(r"\d{4}-(?:0[1-9]|1[0-2])", period):
        raise ValueError("核对月份格式应为 YYYY-MM")
    batch_no = datetime.now().strftime("%Y%m%d%H%M%S") + "-" + hashlib.sha256(raw).hexdigest()[:8]
    seen = set()
    billed_certificates = set()
    findings = []
    with db:
        for line_no, original in rows:
            row = normalized_row(original)
            code = str(row.get("asset_code") or "").strip()
            cert_no = str(row.get("cert_no") or "").strip()
            if cert_no:
                billed_certificates.add(cert_no)
            try:
                charged = amount(row.get("amount"))
                key = (code, cert_no, charged)
                if key in seen:
                    finding, detail = "疑似重复", "本批次存在相同设备、证书和金额；原始行已保留"
                else:
                    finding, detail = reconcile_line(db, code, cert_no, charged)
                seen.add(key)
            except ValueError as exc:
                charged, finding, detail = None, "数据错误", str(exc)
            db.execute("""INSERT INTO billing_record(batch_no,line_no,asset_code,cert_no,amount,raw_data,finding,detail,imported_at)
                VALUES(?,?,?,?,?,?,?,?,?)""", (batch_no, line_no, code, cert_no, charged,
                json.dumps(original, ensure_ascii=False, default=str), finding, detail, now()))
            findings.append({"line_no": line_no, "finding": finding, "detail": detail})
        if period:
            for cert in db.execute("""SELECT c.cert_no,e.asset_code FROM certificate c
                JOIN equipment e ON e.id=c.equipment_id
                WHERE c.state='confirmed' AND c.cal_date LIKE ? AND e.deleted=0""", (period + "-%",)):
                if cert["cert_no"] in billed_certificates:
                    continue
                detail = f"{period} 的已确认校准证书未出现在本批清单"
                db.execute("""INSERT INTO billing_record(batch_no,line_no,asset_code,cert_no,amount,raw_data,finding,detail,imported_at)
                    VALUES(?,?,?,?,?,?,?,?,?)""", (batch_no, 0, cert["asset_code"], cert["cert_no"], None,
                    "{}", "清单缺失", detail, now()))
                findings.append({"line_no": 0, "finding": "清单缺失", "detail": detail, "cert_no": cert["cert_no"]})
        audit(db, "billing_batch", None, "导入", None, {"batch_no": batch_no, "rows": len(rows), "calibration_month": period}, operator)
    return {"batch_no": batch_no, "rows": len(rows), "findings": findings}


def workbook_bytes(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def backup_bytes(db_path=DB_FILE, pdf_dir=PDF_DIR):
    stream = io.BytesIO()
    with BACKUP_LOCK:
        snapshot = db_path.parent / "_backup_snapshot.sqlite3"
        source = connect(db_path)
        try:
            target = sqlite3.connect(snapshot)
            try:
                source.backup(target)
            finally:
                target.close()
        finally:
            source.close()
        try:
            with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.write(snapshot, "register.sqlite3")
                for pdf in pdf_dir.glob("*.pdf") if pdf_dir.exists() else ():
                    archive.write(pdf, f"certificates/{pdf.name}")
        finally:
            snapshot.unlink(missing_ok=True)
    return stream.getvalue()


class Handler(BaseHTTPRequestHandler):
    def respond(self, status, data, content_type="application/json; charset=utf-8", filename=None):
        raw = data if isinstance(data, bytes) else json.dumps(data, ensure_ascii=False, default=str).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        if filename:
            self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        self.end_headers()
        self.wfile.write(raw)

    def do_GET(self):
        path = urlparse(self.path).path
        try:
            if path == "/":
                return self.respond(200, (ROOT / "static" / "index.html").read_bytes(), "text/html; charset=utf-8")
            if path == "/static/app.js":
                return self.respond(200, (ROOT / "static" / "app.js").read_bytes(), "text/javascript; charset=utf-8")
            if path == "/static/style.css":
                return self.respond(200, (ROOT / "static" / "style.css").read_bytes(), "text/css; charset=utf-8")
            with closing(connect()) as db:
                if path == "/api/equipment":
                    query = parse_qs(urlparse(self.path).query)
                    search = query.get("q", [""])[0]
                    rows = db.execute("SELECT * FROM equipment WHERE deleted=0 AND (asset_code LIKE ? OR name LIKE ?) ORDER BY id DESC LIMIT 1000",
                                      (f"%{search}%", f"%{search}%")).fetchall()
                    return self.respond(200, [equipment_view(db, row) for row in rows])
                if path == "/api/certificates":
                    return self.respond(200, [dict(row) for row in db.execute("""SELECT c.*,e.asset_code FROM certificate c
                        LEFT JOIN equipment e ON e.id=c.equipment_id ORDER BY c.id DESC LIMIT 1000""")])
                if re.fullmatch(r"/api/certificates/\d+/file", path):
                    cert_id = int(path.split("/")[3])
                    cert = db.execute("SELECT file_path FROM certificate WHERE id=?", (cert_id,)).fetchone()
                    if not cert:
                        raise ValueError("证书不存在")
                    return self.respond(200, Path(cert["file_path"]).read_bytes(), "application/pdf", f"certificate-{cert_id}.pdf")
                if path == "/api/billing":
                    return self.respond(200, [dict(row) for row in db.execute("SELECT * FROM billing_record ORDER BY id DESC LIMIT 1000")])
                if path == "/api/submissions":
                    return self.respond(200, [dict(row) for row in db.execute("""SELECT s.*,e.asset_code,e.name
                        FROM submission s JOIN equipment e ON e.id=s.equipment_id ORDER BY s.id DESC LIMIT 1000""")])
                if path == "/api/audit":
                    return self.respond(200, [dict(row) for row in db.execute("SELECT * FROM audit_log ORDER BY id DESC LIMIT 1000")])
                if path == "/api/dashboard":
                    equipment = [equipment_view(db, row) for row in db.execute("SELECT * FROM equipment WHERE deleted=0")]
                    today = date.today()
                    next_month = add_months(today.replace(day=1), 1)
                    after_next = add_months(next_month, 1)
                    cal_count = db.execute("SELECT COUNT(DISTINCT equipment_id) FROM certificate WHERE state='confirmed' AND cal_date>=? AND cal_date<?",
                        (today.replace(day=1).isoformat(), add_months(today.replace(day=1), 1).isoformat())).fetchone()[0]
                    sent_count = db.execute("SELECT COUNT(DISTINCT equipment_id) FROM submission WHERE sent_date>=? AND sent_date<?",
                        (today.replace(day=1).isoformat(), add_months(today.replace(day=1), 1).isoformat())).fetchone()[0]
                    in_transit = [dict(row) for row in db.execute("""SELECT s.id,e.asset_code,e.name,s.sent_date
                        FROM submission s JOIN equipment e ON e.id=s.equipment_id
                        WHERE s.returned_date IS NULL AND e.deleted=0 ORDER BY s.sent_date""")]
                    latest_batch = db.execute("SELECT batch_no FROM billing_record ORDER BY id DESC LIMIT 1").fetchone()
                    billing_findings = ([dict(row) for row in db.execute("""SELECT asset_code,cert_no,finding,detail
                        FROM billing_record WHERE batch_no=? AND finding<>'匹配' ORDER BY id LIMIT 100""",
                        (latest_batch["batch_no"],))] if latest_batch else [])
                    return self.respond(200, {"total": len(equipment), "calibrated_this_month": cal_count,
                        "sent_this_month": sent_count, "in_transit": in_transit,
                        "billing_batch": latest_batch["batch_no"] if latest_batch else None,
                        "billing_findings": billing_findings,
                        "statuses": {status: sum(item["status"] == status for item in equipment) for status in ("正常", "超期", "异常", "待校准", "停用")},
                        "next_month": [item for item in equipment if item["next_cal_date"] and next_month.isoformat() <= item["next_cal_date"] < after_next.isoformat()],
                        "abnormal": [item for item in equipment if item["status"] in ("异常", "超期")]})
                if path == "/api/export/equipment":
                    headers = ["资产编号", "设备名称", "型号", "出厂编号", "厂家", "分类", "部门", "地点", "周期(月)", "上次校准", "下次到期", "状态", "异常", "备注"]
                    rows = []
                    for record in db.execute("SELECT * FROM equipment WHERE deleted=0 ORDER BY asset_code"):
                        item = equipment_view(db, record)
                        rows.append([str(item.get(k) or "") for k in ("asset_code", "name", "model", "serial_no", "manufacturer", "category", "department", "location", "calibration_cycle", "last_cal_date", "next_cal_date", "status", "abnormal_flag", "abnormal_note")])
                    return self.respond(200, workbook_bytes(headers, rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "equipment.xlsx")
                if path == "/api/export/audit":
                    rows = list(db.execute("SELECT entity,entity_id,action,old_data,new_data,operator,timestamp FROM audit_log ORDER BY id"))
                    return self.respond(200, workbook_bytes(["实体", "ID", "操作", "原值", "新值", "操作人", "时间"], rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "audit.xlsx")
                if path == "/api/export/billing":
                    rows = list(db.execute("SELECT batch_no,line_no,asset_code,cert_no,amount,finding,detail FROM billing_record ORDER BY id"))
                    return self.respond(200, workbook_bytes(["批次", "行号", "资产编号", "证书编号", "金额", "结果", "详情"], rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "billing.xlsx")
                if path == "/api/backup":
                    return self.respond(200, backup_bytes(), "application/zip", "calibration-backup.zip")
            self.respond(404, {"error": "页面不存在"})
        except (ValueError, FileNotFoundError, sqlite3.Error) as exc:
            self.respond(400, {"error": str(exc)})

    def do_POST(self):
        path = urlparse(self.path).path
        size = int(self.headers.get("Content-Length", "0"))
        if size > MAX_UPLOAD * 2 or size < 0:
            return self.respond(413, {"error": "请求过大"})
        try:
            data = json.loads(self.rfile.read(size))
            operator = str(data.get("operator") or "本机用户").strip()[:64]
            with closing(connect()) as db, db:
                if path == "/api/equipment":
                    result = save_equipment(db, data, operator)
                elif re.fullmatch(r"/api/equipment/\d+", path):
                    result = save_equipment(db, data, operator, int(path.rsplit("/", 1)[1]))
                elif re.fullmatch(r"/api/equipment/\d+/archive", path):
                    result = archive_equipment(db, int(path.split("/")[3]), str(data.get("note") or ""), operator)
                elif path == "/api/import/equipment":
                    result = import_equipment(db, data, operator)
                elif path == "/api/submissions":
                    result = save_submission(db, data, operator)
                elif re.fullmatch(r"/api/submissions/\d+/return", path):
                    result = return_submission(db, int(path.split("/")[3]), data, operator)
                elif path == "/api/certificates":
                    result = upload_pdf(db, data, operator)
                elif re.fullmatch(r"/api/certificates/\d+/review", path):
                    result = review_certificate(db, int(path.split("/")[3]), data, operator)
                elif path == "/api/import/billing":
                    result = import_billing(db, data, operator)
                else:
                    return self.respond(404, {"error": "接口不存在"})
            self.respond(200, result)
        except (ValueError, KeyError, TypeError, sqlite3.IntegrityError) as exc:
            self.respond(400, {"error": str(exc)})


if __name__ == "__main__":
    init_db()
    PDF_DIR.mkdir(parents=True, exist_ok=True)
    server = ThreadingHTTPServer(("127.0.0.1", 8765), Handler)
    print("离线计量台账：http://127.0.0.1:8765", flush=True)
    if "--open" in sys.argv:
        threading.Timer(0.2, lambda: webbrowser.open("http://127.0.0.1:8765")).start()
    server.serve_forever()
